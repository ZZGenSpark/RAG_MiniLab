"""Check heading splits, line joins, and allowlisted policy conversion."""

from pathlib import Path

import pytest
from docx import Document

from config import INGEST_SOURCES, POLICIES_DIR, SOURCE_DIR
from rag.preprocess import (
    paragraphs_to_markdown,
    text_to_markdown,
    write_policies,
    write_policy,
)


def test_docx_rule_paragraph_splits_heading_and_body() -> None:
    """Split `3.1 Requirement. Every email...` into a heading and a body."""
    markdown = paragraphs_to_markdown(["3.1 Requirement. Every email must start with a joke."])
    assert "### 3.1 Requirement" in markdown
    assert "Every email must start with a joke." in markdown
    assert "Requirement. Every" not in markdown


def test_glued_pdf_fragment_splits_purpose_from_body() -> None:
    """Split `1. PurposeThis policy` into a heading and a body."""
    markdown = text_to_markdown("1. PurposeThis policy")
    assert "## 1. Purpose" in markdown
    assert "This policy" in markdown
    assert "PurposeThis" not in markdown


def test_allowlist_has_the_five_ingested_policies() -> None:
    """Leave HR v1 and Preparedness v1 out of preprocessing."""
    names = [path.name for path in INGEST_SOURCES]
    assert len(names) == 5
    assert not any(name.startswith("Doofenshmirtz Evil Inc - HR Policy v1") for name in names)
    assert not any("Preparedness Policy v1" in name for name in names)


def test_excluded_binary_is_rejected(tmp_path: Path) -> None:
    """Refuse to convert a policy that is not on the allowlist."""
    excluded = SOURCE_DIR / "Doofenshmirtz Evil Inc - HR Policy v1.0 1.pdf"
    with pytest.raises(ValueError, match="not an allowlisted policy"):
        write_policy(excluded, tmp_path)


def test_title_body_and_numbered_rule_keep_their_heading_levels() -> None:
    """Keep the document title, a parent section, and a numbered rule distinct."""
    markdown = paragraphs_to_markdown(
        [
            "Doofenshmirtz Evil Incorporated",
            "Time & Usage Policy — Version 1.0",
            "1. Purpose",
            "This policy applies.",
            "5.1 Token Allotment. Every employee is issued 1,000,000 tokens.",
        ]
    )

    assert markdown.startswith("# Time & Usage Policy — Version 1.0\n")
    assert "Doofenshmirtz Evil Incorporated" not in markdown
    assert "## 1. Purpose" in markdown
    assert "This policy applies." in markdown
    assert "### 5.1 Token Allotment" in markdown
    assert "Every employee is issued 1,000,000 tokens." in markdown
    assert "Allotment. Every" not in markdown


def test_rule_without_a_sentence_stays_a_heading() -> None:
    """Leave a numbered rule that has no following sentence as a heading only."""
    markdown = paragraphs_to_markdown(
        [
            "HR Policy — Version 2.0",
            "7.3 Weekend Abandonment",
        ]
    )
    assert "### 7.3 Weekend Abandonment" in markdown
    assert "Weekend Abandonment." not in markdown


def test_pdf_wraps_join_without_merging_a_table() -> None:
    """Join a wrapped sentence, and leave capitalized table rows on their own lines."""
    markdown = text_to_markdown(
        "\n".join(
            [
                "Health Policy — Version 1.0",
                "1. Purpose",
                "Employees receive Focus",
                "Tokens",
                "each Monday.",
                "4. Protein",
                "The table below is a general guideline.",
                "Height Range",
                "Body Weight",
            ]
        )
    )

    assert "Employees receive Focus Tokens each Monday." in markdown
    assert "The table below is a general guideline.\nHeight Range\nBody Weight" in markdown
    assert "guideline. Height" not in markdown
    assert "Height Range Body Weight" not in markdown


def test_glued_version_line_splits_before_the_first_section() -> None:
    """Separate a version number that was glued to the following section heading."""
    markdown = text_to_markdown("Time Policy — Version 1.01. PurposeThis policy applies.")
    assert markdown.startswith("# Time Policy — Version 1.0\n")
    assert "## 1. Purpose" in markdown
    assert "This policy applies." in markdown
    assert "PurposeThis" not in markdown
    assert "Version 1.01" not in markdown


def test_write_policy_names_the_file_from_the_title(tmp_path: Path) -> None:
    """Name the markdown file from the policy title and version."""
    source = tmp_path / "Doofenshmirtz Evil Inc - HR Policy v2.0.docx"
    document = Document()
    document.add_paragraph("HR Policy — Version 2.0")
    document.add_paragraph("7.3 Weekend Abandonment. Food left in the refrigerator is discarded.")
    document.save(str(source))
    dest = tmp_path / "out"
    dest.mkdir()

    written = write_policy(source, dest)

    assert written.name == "hr-policy-v2.0.md"
    text = written.read_text(encoding="utf-8")
    assert text.startswith("# HR Policy — Version 2.0\n")
    assert "### 7.3 Weekend Abandonment" in text
    assert "Food left in the refrigerator is discarded." in text


def test_write_policy_requires_a_title(tmp_path: Path) -> None:
    """Reject a document that never produces a title heading."""
    source = tmp_path / "Doofenshmirtz Evil Inc - HR Policy v2.0.docx"
    document = Document()
    document.add_paragraph("7.3 Weekend Abandonment. Food left in the refrigerator is discarded.")
    document.save(str(source))

    with pytest.raises(ValueError, match="did not produce a policy title"):
        write_policy(source, tmp_path / "out")


def test_write_policy_rejects_a_non_document_suffix(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Reject an allowlisted path that is neither pdf nor docx."""
    monkeypatch.setattr("rag.preprocess.ALLOWED_NAMES", frozenset({"notes.txt"}))
    source = tmp_path / "notes.txt"
    source.write_text("HR Policy — Version 2.0\n", encoding="utf-8")

    with pytest.raises(ValueError, match="not a pdf or docx"):
        write_policy(source, tmp_path / "out")


def test_write_policies_matches_the_checked_in_markdown(tmp_path: Path) -> None:
    """Regenerate the five policies and keep the checked-in markdown unchanged."""
    written = write_policies(dest_dir=tmp_path)
    assert [path.name for path in written] == [
        "hr-policy-v2.0.md",
        "health-and-wellness-policy-v1.0.md",
        "preparedness-policy-v2.0.md",
        "time-and-usage-policy-v2.0.md",
        "time-and-usage-policy-v1.0.md",
    ]
    for path in written:
        expected = POLICIES_DIR / path.name
        assert path.read_text(encoding="utf-8") == expected.read_text(encoding="utf-8")
